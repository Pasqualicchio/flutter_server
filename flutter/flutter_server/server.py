from flask import Flask, request, redirect, url_for, render_template, session, jsonify, send_from_directory, Response, render_template_string, abort
from werkzeug.security import generate_password_hash, check_password_hash
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from functools import wraps
from urllib.parse import unquote
from flask_mail import Mail, Message
import sqlite3
import os
import math
import pandas as pd
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent
app = Flask(__name__)
app.secret_key = 'supersecretkey'

# Cartella di upload coerente
UPLOAD_FOLDER = BASE_DIR / 'uploads'
UPLOAD_FOLDER.mkdir(exist_ok=True)
DB_PATH = BASE_DIR / 'users.db'
LOG_FILE = BASE_DIR / 'upload_log.txt'

# Configura il server SMTP
app.config['MAIL_SERVER'] = 'smtp.gmail.com'
app.config['MAIL_PORT'] = 587
app.config['MAIL_USE_TLS'] = True
app.config['MAIL_USERNAME'] = 'tuo.email@gmail.com'  # Cambia con il tuo indirizzo
app.config['MAIL_PASSWORD'] = 'tua_app_password'      # Cambia con la tua password/app password

mail = Mail(app)

def get_files_info(folder, allowed_exts=None):
    result = []
    for root, _, files in os.walk(folder):
        for file in files:
            if allowed_exts and not file.lower().endswith(tuple(allowed_exts)):
                continue
            full_path = Path(root) / file
            rel_path = full_path.relative_to(folder)
            stat = full_path.stat()
            ext = file.lower().split('.')[-1]
            file_type = 'image' if ext in ['jpg', 'jpeg', 'png'] else \
                        'video' if ext in ['mp4', 'avi'] else \
                        'pdf' if ext == 'pdf' else \
                        'audio' if ext in ['mp3', 'wav'] else 'other'
            result.append({
                'name': file,
                'path': str(rel_path).replace("\\", "/"),
                'type': file_type,
                'size_kb': round(stat.st_size / 1024, 1),
                'created': datetime.fromtimestamp(stat.st_ctime).strftime("%Y-%m-%d %H:%M"),
                'modified': datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M")
            })
    return result

def paginate(items, page=1, per_page=20):
    total = len(items)
    total_pages = math.ceil(total / per_page)
    start = (page - 1) * per_page
    end = start + per_page
    return items[start:end], total_pages

def init_db():
    with sqlite3.connect(DB_PATH) as conn:
        c = conn.cursor()
        c.execute('''CREATE TABLE IF NOT EXISTS users (
                        id INTEGER PRIMARY KEY AUTOINCREMENT,
                        username TEXT UNIQUE NOT NULL,
                        password TEXT NOT NULL
                    )''')
        conn.commit()

def current_user():
    return session.get('user')

def login_required(view_func):
    @wraps(view_func)
    def wrapped(*args, **kwargs):
        if not current_user():
            return redirect(url_for('login'))
        return view_func(*args, **kwargs)
    return wrapped

def safe_name(name):
    return name.replace("=", "").replace("\\", "_").replace("/", "_").replace(" ", "_").strip()

def update_excel_file(macro_folder, macro_name, log_data, image_path):
    excel_path = Path(macro_folder) / f"log_{safe_name(macro_name)}.xlsx"
    headers = ["Timestamp", "Macro", "Linea", "Commessa", "Matricola", "Item", "Filename", "Preview"]

    if not excel_path.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "Log"
        ws.append(headers)
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="4F81BD", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")
            cell.border = Border(left=Side(style='thin'), right=Side(style='thin'),
                                 top=Side(style='thin'), bottom=Side(style='thin'))
    else:
        wb = load_workbook(excel_path)
        ws = wb.active

    row_idx = ws.max_row + 1
    for col_idx, value in enumerate(log_data, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = Alignment(horizontal="center")
        cell.border = Border(left=Side(style='thin'), right=Side(style='thin'),
                             top=Side(style='thin'), bottom=Side(style='thin'))

    image_path = Path(image_path)
    if image_path.exists() and image_path.suffix.lower() in ['.jpg', '.jpeg', '.png']:
        try:
            img = XLImage(str(image_path))
            img.width = 100
            img.height = 70
            ws.add_image(img, f"H{row_idx}")
            ws.row_dimensions[row_idx].height = 55
        except Exception as e:
            print(f"[Errore immagine Excel] {e}")

    for col in range(1, ws.max_column + 1):
        max_len = max(len(str(ws.cell(row=row, column=col).value or '')) for row in range(1, ws.max_row + 1))
        ws.column_dimensions[get_column_letter(col)].width = max(max_len + 2, 15)

    wb.save(excel_path)

def generate_global_excel():
    global_path = BASE_DIR / 'global_log.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.title = "Global Log"
    headers = ["Timestamp", "Macro", "Linea", "Commessa", "Matricola", "Item", "Filename"]
    ws.append(headers)

    for root, _, files in os.walk(UPLOAD_FOLDER):
        for f in files:
            if f.lower().startswith("log_") and f.lower().endswith(".xlsx"):
                try:
                    path = os.path.join(root, f)
                    local_wb = load_workbook(path)
                    local_ws = local_wb.active
                    for row in local_ws.iter_rows(min_row=2, max_col=7, values_only=True):
                        ws.append(row)
                except Exception as e:
                    print(f"[Errore lettura {f}] {e}")

    wb.save(global_path)

@app.route('/send-email', methods=['POST'])
def send_email():
    data = request.get_json()
    email = data.get('email')
    message = data.get('message', '')
    files = data.get('files', [])

    if not email or not files:
        return jsonify(success=False, message="Dati incompleti"), 400

    try:
        msg = Message("📎 File condivisi", sender=app.config['MAIL_USERNAME'], recipients=[email])
        msg.body = message or "Ecco i file selezionati."

        for filepath in files:
            filename = os.path.basename(filepath)
            full_path = UPLOAD_FOLDER / filename
            if full_path.exists():
                with open(full_path, 'rb') as f:
                    msg.attach(filename, "application/octet-stream", f.read())

        mail.send(msg)
        return jsonify(success=True)

    except Exception as e:
        print("Errore email:", e)
        return jsonify(success=False, message=str(e)), 500

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password')

        with sqlite3.connect(DB_PATH) as conn:
            user = conn.execute("SELECT * FROM users WHERE username = ?", (username,)).fetchone()
            if user and check_password_hash(user[2], password):
                session['user'] = username
                return redirect(url_for('browse_ui'))

        return render_template('login.html', error="Credenziali errate.")

    return render_template('login.html')

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

@app.route('/browse-ui')
@login_required
def browse_ui():
    images_info = get_files_info(UPLOAD_FOLDER)
    excels = get_files_info(UPLOAD_FOLDER, allowed_exts=['.xlsx'])
    page = 1
    per_page = 20
    paginated_images, total_pages = paginate(images_info, page=page, per_page=per_page)
    return render_template('browse.html', images=paginated_images, excels=excels, page=page, total_pages=total_pages)

@app.route('/upload-ui')
@login_required
def upload_ui():
    return render_template('upload_ui.html')

@app.route('/upload-advanced-ui')
@login_required
def upload_advanced_ui():
    return render_template('upload_advanced_ui.html')

@app.route('/download/<path:filepath>')
@login_required
def download(filepath):
    return send_from_directory(UPLOAD_FOLDER, filepath, as_attachment=True)

@app.route('/preview-excel/<path:filepath>')
@login_required
def preview_excel(filepath):
    return f"Anteprima non disponibile per: {filepath}"

@app.route('/delete-file/<path:filepath>', methods=['DELETE'])
@login_required
def delete_file(filepath):
    try:
        full_path = UPLOAD_FOLDER / Path(filepath)
        if full_path.exists():
            full_path.unlink()
            return jsonify(status="success")
        else:
            return jsonify(status="error", message="File non trovato"), 404
    except Exception as e:
        return jsonify(status="error", message=str(e)), 500

@app.route('/upload', methods=['POST'])
@login_required
def upload():
    try:
        macro = request.form.get('macro')
        linea = request.form.get('linea', '')
        commessa = request.form.get('commessa')
        matricola = request.form.get('matricola')
        item = request.form.get('item')
        images = request.files.getlist('image')
        videos = request.files.getlist('video')

        if not all([macro, commessa, matricola, item]) or (not images and not videos):
            return jsonify({'status': 'error', 'message': 'Dati mancanti o nessun file'}), 400

        macro_folder = os.path.join(UPLOAD_FOLDER, safe_name(macro))
        os.makedirs(macro_folder, exist_ok=True)

        photo_folder = os.path.join(
            macro_folder,
            safe_name(commessa),
            safe_name(matricola),
            safe_name(item)
        )
        os.makedirs(photo_folder, exist_ok=True)

        uploaded_files = []

        for image in images:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S_%f')
            filename = f"{timestamp}_{safe_name(image.filename)}"
            image_path = os.path.join(photo_folder, filename)
            image.save(image_path)

            with open(LOG_FILE, 'a', encoding='utf-8') as log:
                log.write(f"{datetime.now().isoformat()} | {macro} | {linea} | {commessa} | {matricola} | {item} | {filename}\n")

            update_excel_file(macro_folder, macro, [
                datetime.now().isoformat(), macro, linea, commessa, matricola, item, filename
            ], image_path)

            uploaded_files.append(filename)

        for video in videos:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S_%f')
            filename = f"{timestamp}_{safe_name(video.filename)}"
            video_path = os.path.join(photo_folder, filename)
            video.save(video_path)

            with open(LOG_FILE, 'a', encoding='utf-8') as log:
                log.write(f"{datetime.now().isoformat()} | {macro} | {linea} | {commessa} | {matricola} | {item} | {filename}\n")

            uploaded_files.append(filename)

        return jsonify({'status': 'success', 'filenames': uploaded_files}), 200

    except Exception as e:
        print(f"[❌ Errore upload] {e}")
        return jsonify({'status': 'error', 'message': str(e)}), 500
    
if __name__ == '__main__':
    app.run(debug=True)

